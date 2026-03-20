#!/usr/bin/env python3
"""Simple Markdown to DOCX converter for project docs."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def convert(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    in_code = False

    for raw in lines:
        line = raw.rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            doc.add_paragraph(line)
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        elif stripped.startswith(tuple(f"{n}. " for n in range(1, 10))):
            # Preserve numbered list text in plain paragraph for compatibility.
            doc.add_paragraph(stripped)
        else:
            doc.add_paragraph(line)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    md_path = Path(args.input).resolve()
    out_path = Path(args.output).resolve()
    convert(md_path, out_path)
    print(f"[Done] DOCX generated: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

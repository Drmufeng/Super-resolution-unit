#!/usr/bin/env python3
"""Run one-click evaluation and generate CSV/Markdown/optional DOCX report."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parent
SR30 = ROOT / "sr30_workspace"
SRCNN = ROOT / "srcnn"
REPORT_DIR = ROOT / "reports"

AVG_PSNR_RE = re.compile(r"\[Done\]\s+avg_psnr=([0-9eE+\-.]+)")
AVG_SSIM_RE = re.compile(r"\[Done\]\s+avg_ssim=([0-9eE+\-.]+)")
COUNT_RE = re.compile(r"\[Done\]\s+count=(\d+)")


def parse_metrics(text: str) -> tuple[int | None, float | None, float | None]:
    c = COUNT_RE.search(text)
    p = AVG_PSNR_RE.search(text)
    s = AVG_SSIM_RE.search(text)
    return (int(c.group(1)) if c else None, float(p.group(1)) if p else None, float(s.group(1)) if s else None)


def run_eval(workspace: Path, script_name: str, config_path: Path) -> dict:
    cmd = [sys.executable, str(workspace / script_name), "--config", str(config_path)]
    print("[Cmd]", " ".join(cmd), flush=True)
    started = dt.datetime.now()
    proc = subprocess.run(cmd, cwd=str(workspace), capture_output=True, text=True)
    elapsed = (dt.datetime.now() - started).total_seconds()
    out = (proc.stdout or "") + ("\n" + proc.stderr if proc.stderr else "")
    count, psnr, ssim = parse_metrics(out)
    status = "ok" if proc.returncode == 0 and psnr is not None and ssim is not None else "failed"
    return {
        "workspace": workspace.name,
        "config": str(config_path.relative_to(workspace)),
        "status": status,
        "returncode": proc.returncode,
        "count": count,
        "avg_psnr": psnr,
        "avg_ssim": ssim,
        "seconds": round(elapsed, 2),
        "message": "" if status == "ok" else (out.strip().splitlines()[-1] if out.strip() else "unknown error"),
    }


def build_jobs(mode: str) -> list[tuple[Path, str, Path]]:
    jobs: list[tuple[Path, str, Path]] = []
    if mode in {"all", "sr30"}:
        for cfg in [
            SR30 / "configs" / "eval_x2_set14_edsr.json",
            SR30 / "configs" / "eval_x2_b100_edsr.json",
            SR30 / "configs" / "eval_x2_urban100_edsr.json",
            SR30 / "configs" / "eval_x4_set14_edsr.json",
            SR30 / "configs" / "eval_x4_b100_edsr.json",
            SR30 / "configs" / "eval_x4_urban100_edsr.json",
        ]:
            jobs.append((SR30, "eval_psnr_edsr.py", cfg))

    if mode in {"all", "srcnn"}:
        for cfg in [
            SRCNN / "configs" / "eval_x2_set14.json",
            SRCNN / "configs" / "eval_x2_b100.json",
            SRCNN / "configs" / "eval_x2_urban100.json",
            SRCNN / "configs" / "eval_x4_set14.json",
            SRCNN / "configs" / "eval_x4_b100.json",
            SRCNN / "configs" / "eval_x4_urban100.json",
        ]:
            jobs.append((SRCNN, "eval.py", cfg))
    return jobs


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["workspace", "config", "status", "returncode", "count", "avg_psnr", "avg_ssim", "seconds", "message"],
        )
        writer.writeheader()
        for r in rows:
            row = dict(r)
            if isinstance(row.get("avg_psnr"), float):
                row["avg_psnr"] = f"{row['avg_psnr']:.4f}"
            if isinstance(row.get("avg_ssim"), float):
                row["avg_ssim"] = f"{row['avg_ssim']:.4f}"
            writer.writerow(row)


def write_markdown(rows: list[dict], path: Path) -> None:
    ok = [r for r in rows if r["status"] == "ok"]
    failed = [r for r in rows if r["status"] != "ok"]
    lines = [
        "# 一键测试报告",
        "",
        f"- 生成时间: {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- 总任务: {len(rows)}",
        f"- 成功: {len(ok)}",
        f"- 失败: {len(failed)}",
        "",
        "| workspace | config | status | count | avg_psnr | avg_ssim | seconds |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for r in rows:
        psnr = "" if r["avg_psnr"] is None else f"{r['avg_psnr']:.4f}"
        ssim = "" if r["avg_ssim"] is None else f"{r['avg_ssim']:.4f}"
        count = "" if r["count"] is None else str(r["count"])
        lines.append(
            f"| {r['workspace']} | {r['config']} | {r['status']} | {count} | {psnr} | {ssim} | {r['seconds']:.2f} |"
        )

    if failed:
        lines += ["", "## 失败详情", ""]
        for r in failed:
            lines.append(f"- `{r['workspace']} / {r['config']}`: {r['message']}")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def try_write_docx(md_path: Path, docx_path: Path) -> bool:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return False

    doc = Document()
    doc.add_heading("一键测试报告", level=1)
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            continue
        doc.add_paragraph(line)
    doc.save(docx_path)
    return True


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["all", "sr30", "srcnn"], default="all")
    parser.add_argument("--with-docx", action="store_true")
    args = parser.parse_args()

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    jobs = build_jobs(args.mode)
    rows = []
    for workspace, script_name, cfg in jobs:
        if not cfg.is_file():
            rows.append(
                {
                    "workspace": workspace.name,
                    "config": str(cfg),
                    "status": "failed",
                    "returncode": -1,
                    "count": None,
                    "avg_psnr": None,
                    "avg_ssim": None,
                    "seconds": 0.0,
                    "message": "config missing",
                }
            )
            continue
        rows.append(run_eval(workspace, script_name, cfg))

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path = REPORT_DIR / f"test_report_{stamp}.csv"
    md_path = REPORT_DIR / f"test_report_{stamp}.md"
    write_csv(rows, csv_path)
    write_markdown(rows, md_path)
    shutil.copy2(csv_path, REPORT_DIR / "test_report_latest.csv")
    shutil.copy2(md_path, REPORT_DIR / "test_report_latest.md")

    print(f"[Done] CSV: {csv_path}")
    print(f"[Done] MD: {md_path}")

    if args.with_docx:
        docx_path = REPORT_DIR / f"test_report_{stamp}.docx"
        if try_write_docx(md_path, docx_path):
            shutil.copy2(docx_path, REPORT_DIR / "test_report_latest.docx")
            print(f"[Done] DOCX: {docx_path}")
        else:
            print("[Warn] python-docx not installed, skip docx output")

    failed = sum(1 for r in rows if r["status"] != "ok")
    print(f"[Done] total={len(rows)}, failed={failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
